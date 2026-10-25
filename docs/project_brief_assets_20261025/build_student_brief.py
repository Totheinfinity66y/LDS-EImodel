from pathlib import Path

from pptx import Presentation
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE
from pptx.enum.text import MSO_ANCHOR, PP_ALIGN
from pptx.oxml.xmlchemy import OxmlElement
from pptx.oxml.ns import qn
from pptx.util import Inches, Pt


ROOT = Path(__file__).resolve().parents[2]
ASSETS = Path(__file__).resolve().parent
OUT = ROOT / "docs" / "项目进展简报_交大红色科学期刊风.pptx"

W, H = 13.333, 7.5
BG = "F8F7F2"
RED = "B21E35"
RED2 = "D34A4C"
INK = "282927"
MID = "5C5D59"
MUTED = "878781"
PALE = "E9E7E0"
PALE_RED = "F1DEDD"
WHITE = "FFFFFF"
FONT = "Microsoft YaHei"
MATH = "Cambria Math"


def rgb(hex_value):
    return RGBColor.from_string(hex_value)


def set_run(run, size=16, color=INK, bold=False, font=FONT, italic=False):
    run.font.name = font
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.italic = italic
    run.font.color.rgb = rgb(color)
    rpr = run._r.get_or_add_rPr()
    rfonts = rpr.find(qn("a:rFonts"))
    if rfonts is None:
        rfonts = OxmlElement("a:rFonts")
        rpr.insert(0, rfonts)
    rfonts.set(qn("a:latin"), font)
    rfonts.set(qn("a:ea"), font)


def textbox(slide, x, y, w, h, text, size=16, color=INK, bold=False,
            font=FONT, align=PP_ALIGN.LEFT, valign=MSO_ANCHOR.TOP,
            margin=0, italic=False, leading=1.05):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    p = tf.paragraphs[0]
    p.alignment = align
    p.space_before = Pt(0)
    p.space_after = Pt(0)
    p.line_spacing = leading
    r = p.add_run()
    r.text = text
    set_run(r, size, color, bold, font, italic)
    return shape


def rich_lines(slide, x, y, w, h, lines, align=PP_ALIGN.LEFT,
               valign=MSO_ANCHOR.TOP, margin=0.02, leading=1.0):
    shape = slide.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h))
    tf = shape.text_frame
    tf.clear()
    tf.word_wrap = True
    tf.margin_left = tf.margin_right = Inches(margin)
    tf.margin_top = tf.margin_bottom = Inches(margin)
    tf.vertical_anchor = valign
    for i, spec in enumerate(lines):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.alignment = spec.get("align", align)
        p.space_before = Pt(spec.get("before", 0))
        p.space_after = Pt(spec.get("after", 4))
        p.line_spacing = spec.get("leading", leading)
        r = p.add_run()
        r.text = spec["text"]
        set_run(r, spec.get("size", 16), spec.get("color", INK),
                spec.get("bold", False), spec.get("font", FONT),
                spec.get("italic", False))
    return shape


def rect(slide, x, y, w, h, fill, line=None, radius=False):
    shape = slide.shapes.add_shape(
        MSO_SHAPE.ROUNDED_RECTANGLE if radius else MSO_SHAPE.RECTANGLE,
        Inches(x), Inches(y), Inches(w), Inches(h)
    )
    shape.fill.solid()
    shape.fill.fore_color.rgb = rgb(fill)
    shape.line.fill.background() if line is None else None
    if line:
        shape.line.color.rgb = rgb(line)
        shape.line.width = Pt(0.8)
    return shape


def add_image(slide, filename, x, y, w, h):
    path = ASSETS / filename
    from PIL import Image
    with Image.open(path) as im:
        iw, ih = im.size
    scale = min(w / iw, h / ih)
    rw, rh = iw * scale, ih * scale
    return slide.shapes.add_picture(
        str(path), Inches(x + (w-rw)/2), Inches(y + (h-rh)/2),
        width=Inches(rw), height=Inches(rh)
    )


def new_slide(prs):
    slide = prs.slides.add_slide(prs.slide_layouts[6])
    slide.background.fill.solid()
    slide.background.fill.fore_color.rgb = rgb(BG)
    return slide


def header(slide, section, title, page):
    rect(slide, 0.0, 0.0, W, 0.09, RED)
    textbox(slide, 0.72, 0.34, 3.2, 0.24, section.upper(), 10, RED, True,
            leading=1.0)
    textbox(slide, 0.72, 0.68, 11.85, 0.58, title, 27, INK, True,
            leading=0.98)
    textbox(slide, 12.15, 7.08, 0.48, 0.2, f"{page:02d}", 9, MUTED,
            align=PP_ALIGN.RIGHT)


def footer(slide, text):
    textbox(slide, 0.74, 7.04, 10.8, 0.22, text, 8.5, MUTED,
            leading=1.0)


def formula_band(slide, x, y, w, h, text, size=18):
    rect(slide, x, y, w, h, PALE_RED)
    textbox(slide, x + 0.16, y + 0.08, w - 0.32, h - 0.16, text,
            size, RED, True, font=MATH, valign=MSO_ANCHOR.MIDDLE,
            align=PP_ALIGN.CENTER, leading=1.0)


def image_slide(prs, section, title, image, page, kicker=None):
    slide = new_slide(prs)
    header(slide, section, title, page)
    add_image(slide, image, 0.55, 1.42, 12.23, 5.38)
    if kicker:
        textbox(slide, 0.78, 6.49, 11.5, 0.35, kicker, 15, RED, True,
                align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    return slide


def build():
    prs = Presentation()
    prs.slide_width = Inches(W)
    prs.slide_height = Inches(H)
    prs.core_properties.title = "项目进展：Handcrafted E/I Decompression v1"
    prs.core_properties.subject = "面向零背景同学的项目汇报"
    prs.core_properties.author = ""
    prs.core_properties.keywords = "E/I, L/D/S, disinhibition, response-state"

    # 01 — Cover
    s = new_slide(prs)
    rect(s, 0.78, 0.78, 0.08, 1.08, RED)
    textbox(s, 1.03, 0.77, 4.3, 0.28, "项目进展｜HANDCRAFTED E/I V1", 11,
            RED, True)
    textbox(s, 1.03, 1.35, 4.9, 1.42, "短简规则\n如何变成连续行动？",
            33, INK, True, leading=0.98)
    textbox(s, 1.05, 3.13, 4.1, 0.68,
            "固定连接的兴奋／抑制网络\n从已知 L / D / S 到四次位置读出",
            17, MID, leading=1.12)
    formula_band(s, 1.04, 4.36, 3.8, 0.68, "L  ·  D  ·  S  →  R₁  →  R₂  →  R₃  →  R₄", 16)
    add_image(s, "slide01_cover_v2.png", 5.45, 0.58, 7.42, 6.28)
    footer(s, "15 分钟项目介绍  ·  概念插画用于解释结构，不代表实测神经影像")

    # 02 — Task definition
    s = new_slide(prs)
    header(s, "01 / 任务定义", "六个位置，一条四步序列", 2)
    add_image(s, "slide02_task_v2.png", 0.55, 1.48, 6.4, 4.92)
    textbox(s, 7.18, 1.56, 5.25, 0.62, "输入只有三个量", 21, INK, True)
    rich_lines(s, 7.2, 2.28, 5.2, 1.32, [
        {"text":"L  起点  0–5", "size":17, "color":MID},
        {"text":"D  方向  +1 / −1", "size":17, "color":MID},
        {"text":"S  中间步长  1 / 2", "size":17, "color":MID},
    ], leading=1.0)
    textbox(s, 7.18, 3.68, 5.3, 0.38, "只在 Z₆ 圆环上移动", 15, MID)
    formula_band(s, 7.13, 4.18, 5.4, 1.2,
                 "R₁ = L\nR₂ = L + D\nR₃ = L + D + DS\nR₄ = L + 2D + DS  (mod 6)", 16)
    textbox(s, 7.2, 5.68, 4.9, 0.32, "例： (0, +1, 2)  →  0 → 1 → 3 → 4", 16,
            RED, True, font=MATH)
    footer(s, "24 种条件 = 6 个起点 × 2 个方向 × 2 个步长；模型只接收 L / D / S 输入。")

    # 03 — Scientific questions
    s = new_slide(prs)
    header(s, "02 / 科学问题", "输出正确之外，还要追踪信息怎样走", 3)
    add_image(s, "slide03_question_v2.png", 0.7, 1.48, 11.95, 3.55)
    labels = [
        (0.82, "保持", "输入撤除后\n起点与规则还在吗？"),
        (3.25, "选择", "何时打开\n正确通路？"),
        (6.0, "继承", "用过的规则\n怎样由响应承接？"),
        (9.12, "读出", "每次先读位置\n再开始下一步"),
    ]
    for x, title, body in labels:
        textbox(s, x, 5.05, 2.0, 0.32, title, 17, RED, True,
                align=PP_ALIGN.CENTER)
        textbox(s, x - 0.05, 5.47, 2.1, 0.63, body, 13.5, MID,
                align=PP_ALIGN.CENTER, leading=1.0)
    footer(s, "v1 直接提供已压缩的 L / D / S；本项目聚焦从规则保持到序列展开的机制。")

    # 04 — E/I network
    s = new_slide(prs)
    header(s, "03 / 网络结构", "固定连接把规则变成可检查的神经活动", 4)
    add_image(s, "slide04_network_v2.png", 0.58, 1.52, 7.38, 4.65)
    textbox(s, 8.08, 1.56, 4.55, 0.38, "四个响应阶段", 20, INK, True)
    textbox(s, 8.08, 2.1, 4.5, 0.8,
            "R₁ 保持起点\nR₂ 形成方向分支\nR₃ 形成方向 × 步长分支\nR₄ 沿分支继续推进",
            15, MID, leading=1.03)
    formula_band(s, 8.05, 4.2, 4.55, 0.75, "76 E  +  28 I  =  104 个率状态", 17)
    textbox(s, 8.1, 5.1, 4.55, 0.72,
            "E 输出兴奋；I 输出抑制。\nI→I 也保持抑制符号。权重手工设定，不做训练。",
            13.5, MID, leading=1.02)
    formula_band(s, 0.9, 6.2, 11.5, 0.55,
                 "τᵢ · drᵢ/dt = −rᵢ + φ( ΣE Wᵢⱼrⱼ − ΣI Wᵢⱼrⱼ + Iᵢ − θᵢ )", 15)
    footer(s, "率为归一化无量纲活动；单元数与参数描述本实现，不主张为生物最小结构。")

    # 05 — Disinhibition
    s = new_slide(prs)
    header(s, "04 / 通路门控", "抑制门默认关闭，释放单元打开目标通路", 5)
    add_image(s, "slide05_disinhibition_v2.png", 0.58, 1.6, 7.45, 4.76)
    textbox(s, 8.2, 1.85, 4.3, 0.42, "两层抑制", 21, INK, True)
    rich_lines(s, 8.2, 2.48, 4.3, 2.0, [
        {"text":"I_gate 压住下游目标", "size":16, "color":MID, "after":10},
        {"text":"I_release 抑制 I_gate", "size":16, "color":MID, "after":10},
        {"text":"目标通路解除抑制后\n固定位置投射才得以传递", "size":16, "color":MID},
    ], leading=1.0)
    formula_band(s, 7.95, 4.92, 4.8, 1.05,
                 "M(E) → I_release ⊣ I_gate ⊣ E_target", 18)
    textbox(s, 8.15, 6.1, 4.5, 0.32, "规则控制通路；内容来自源位置活动", 13.5,
            RED, True)
    footer(s, "概念插画辅助解释；I_release 与 I_gate 在模型中都有真实率状态。")

    # 06 — State inheritance and readout-first
    s = new_slide(prs)
    header(s, "05 / 响应状态", "已选分支承接规则，后续不重复查询", 6)
    add_image(s, "slide06_inheritance_v2.png", 0.62, 1.48, 12.05, 3.58)
    stages = [
        (0.9, "R₁ → R₂", "D 选方向"),
        (3.88, "R₂ → R₃", "活动继承 D；S 选步长"),
        (7.15, "R₃ → R₄", "沿固定分支推进"),
    ]
    for x, title, desc in stages:
        textbox(s, x, 5.08, 2.8, 0.32, title, 17, RED, True,
                align=PP_ALIGN.CENTER, font=MATH)
        textbox(s, x, 5.48, 2.8, 0.45, desc, 13.5, MID,
                align=PP_ALIGN.CENTER)
    formula_band(s, 10.15, 5.16, 2.3, 0.8, "1 → 2 → 4 → 4", 17)
    textbox(s, 1.0, 6.22, 11.2, 0.38,
            "读 R₁ → 转移 → 读 R₂ → 转移 → 读 R₃ → 转移 → 读 R₄",
            16, INK, True, align=PP_ALIGN.CENTER, font=MATH)
    footer(s, "1/2/4/4 是潜在分支数；clock 只提供阶段时间，不提供位置答案。")

    # 07 — Observed validation
    s = new_slide(prs)
    header(s, "06 / v1 验证", "定义条件内，24 条完整序列全部正确", 7)
    add_image(s, "v1_validation_trajectories.png", 0.48, 1.5, 8.25, 5.2)
    textbox(s, 9.0, 1.62, 3.35, 0.36, "无噪声完整序列", 15, MID, True)
    textbox(s, 8.92, 2.0, 3.5, 1.02, "24 / 24", 35, RED, True,
            font=MATH)
    textbox(s, 9.02, 3.06, 3.2, 0.34, "每阶段准确率 100%", 16, INK, True)
    rect(s, 9.0, 3.64, 3.45, 0.018, PALE)
    textbox(s, 9.02, 3.88, 3.3, 0.3, "自动化测试", 13, MID, True)
    textbox(s, 9.02, 4.2, 3.3, 0.54, "16 / 16 通过", 20, INK, True)
    textbox(s, 9.02, 4.95, 3.3, 0.3, "保存轨迹独立复算", 13, MID, True)
    textbox(s, 9.02, 5.28, 3.5, 0.54, "39 实验 · 936 trial", 17, INK, True)
    footer(s, "原始图与复算记录来自 v1 归档；936 个 trial 横跨不同设置与干预，不是独立样本数。")

    # 08 — Causal interventions
    s = new_slide(prs)
    header(s, "07 / 因果干预", "切断时机不同，失效阶段也不同", 8)
    add_image(s, "slide08_interventions_v2.png", 0.82, 1.45, 11.7, 2.28)
    # Editable stage-by-stage matrix: each mark reports full-stage accuracy.
    cols = [0.95, 3.48, 5.73, 7.98, 10.23]
    labels = ["干预", "R₁", "R₂", "R₃", "R₄"]
    for x, label in zip(cols, labels):
        textbox(s, x, 3.82, 2.05, 0.3, label, 12.5,
                RED if x == cols[0] else MID, True,
                align=PP_ALIGN.CENTER, font=MATH)
    rows = [
        ("A  阻断 D release", [100, 0, 0, 0]),
        ("B  读出 R₂ 后切断 M_D", [100, 100, 100, 100]),
        ("C  阻断 S release", [100, 100, 0, 0]),
        ("D  读出 R₃ 后切断 M_D/M_S", [100, 100, 100, 100]),
    ]
    for ri, (label, vals) in enumerate(rows):
        y = 4.23 + ri * 0.47
        textbox(s, cols[0], y, 2.55, 0.31, label, 12.3, INK,
                align=PP_ALIGN.LEFT)
        for ci, val in enumerate(vals):
            x = cols[ci + 1]
            rect(s, x + 0.2, y + 0.04, 1.65, 0.22,
                 RED if val == 100 else PALE)
            textbox(s, x + 0.2, y - 0.015, 1.65, 0.28,
                    "✓" if val == 100 else "—", 11,
                    WHITE if val == 100 else MUTED, True,
                    align=PP_ALIGN.CENTER, valign=MSO_ANCHOR.MIDDLE)
    textbox(s, 1.0, 6.18, 11.2, 0.3,
            "A、C：完整序列 0/24     B、D：完整序列 24/24",
            15, RED, True, align=PP_ALIGN.CENTER)
    footer(s, "A–D 均有等剂量随机 lesion 对照（seed 7/11/19）；对照结果不用于普遍显著性推断。")

    # 09 — Meaning and limits
    s = new_slide(prs)
    header(s, "08 / 结果解释", "这是一项可检验的机制原型证据", 9)
    add_image(s, "slide09_evidence_limits_v2.png", 2.05, 1.46, 9.25, 2.45)
    rect(s, 0.92, 4.1, 0.06, 1.78, RED)
    textbox(s, 1.18, 4.14, 5.0, 0.38, "结果支持", 19, RED, True)
    textbox(s, 1.18, 4.62, 4.95, 1.15,
            "在当前固定网络与所测条件中，\n记忆、去抑制和状态继承\n足以生成定义好的四位置序列。",
            15.5, INK, leading=1.06)
    rect(s, 6.95, 4.1, 0.06, 1.78, PALE)
    textbox(s, 7.22, 4.14, 5.0, 0.38, "尚未证明", 19, MID, True)
    textbox(s, 7.22, 4.62, 5.0, 1.15,
            "大脑采用相同结构、连续吸引子、\n涌现正交子空间，或普适鲁棒性。\n阶段推进目前由外部 clock 控制。",
            15.5, MID, leading=1.06)
    footer(s, "解释范围限于已定义任务、固定拓扑和归档参数。")

    # 10 — Next steps
    s = new_slide(prs)
    header(s, "09 / 下一步", "先测机制边界，再问哪些结构不可少", 10)
    add_image(s, "slide10_roadmap_v2.png", 0.7, 1.52, 11.95, 3.15)
    cards = [
        (0.9, "01  扰动与保持", "持续噪声\n更长 delay\n活动恢复范围"),
        (4.8, "02  必要结构", "简化单元与通路\n联合参数扫描\n量化失效边界"),
        (8.7, "03  自主时序", "内部阶段转换\n单独定义新规格\n保留信息边界"),
    ]
    for x, title, body in cards:
        textbox(s, x, 4.86, 3.3, 0.32, title, 15.5, RED, True)
        textbox(s, x, 5.3, 3.35, 1.0, body, 13.5, MID, leading=1.08)
    footer(s, "从当前 v1 扩展前，先独立定义实验协议与通过判据；学习 L/D/S 属于新版本任务。")

    OUT.parent.mkdir(parents=True, exist_ok=True)
    prs.save(OUT)
    print(OUT)
    print(f"slides={len(prs.slides)} size={W}x{H}")


if __name__ == "__main__":
    build()
